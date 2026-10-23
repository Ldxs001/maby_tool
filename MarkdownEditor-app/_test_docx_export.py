# -*- coding: utf-8 -*-
# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 wUwproject. Licensed under GNU AGPL v3 (or later).
"""docx 导出 + 一键复制 HTML 导出 回归测试（v2.1.x 新增）。

覆盖：
- build_standalone_html：悬浮「一键复制」按钮 + __inlineClone + ClipboardItem 内联逻辑
- docx_export.build_docx：标题/段落(内联样式)/列表/引用/代码块/表格/公式(data-tex)/图片
  - 断言 word/document.xml 含 m:oMath 原生可编辑公式节点
  - 断言 python-docx 可打开、含对应段落与表格、含嵌入图片
- 轻量回归：md_server 导入无错误，既有关键函数仍在
"""
import io
import os
import sys
import tempfile
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import md_server  # noqa: E402  # 先导入，再确保依赖
assert md_server._ensure_docx_deps(), "docx 依赖不可用（应内置于 MarkdownEditor-app/docxdeps/）"
import docx_export  # noqa: E402  # 依赖就绪后再导入
from docx import Document  # noqa: E402
import lxml.etree as ET  # noqa: E402

M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

RESULTS = []


def check(name, ok, extra=""):
    RESULTS.append((name, ok))
    print("%s %s%s" % ("PASS" if ok else "FAIL", name,
                       (" | " + extra) if extra else ""))


def _make_png():
    # add_picture 依赖 Pillow，用其生成一张合法 PNG 作图片嵌入用例
    from PIL import Image
    d = tempfile.mkdtemp(prefix="md_docx_")
    p = os.path.join(d, "px.png")
    Image.new("RGB", (8, 8), (200, 60, 60)).save(p)
    return p


def sample_html(img_path):
    ip = img_path.replace("\\", "/")
    return ('<div class="markdown-body">'
            '<h1>标题一</h1>'
            '<h2>小标题</h2>'
            '<p style="text-align:center;color:#cc0000;font-size:18px">'
            '红色居中 <strong>粗体</strong> 与 <em>斜体</em> 文本。</p>'
            '<ul><li>列表项一</li><li>列表项二</li></ul>'
            '<blockquote style="background-color:#f5f7ff">这是一段引用文字。</blockquote>'
            '<pre><code>x = 1\nprint(x)</code></pre>'
            '<table><tr><th>列A</th><th>列B</th></tr><tr><td>1</td><td>2</td></tr></table>'
            '<div class="md-math-block" data-tex="E = mc^2"></div>'
            '<p>行内公式 <span class="md-math-inline" data-tex="a^2+b^2=c^2"></span> 结束。</p>'
            '<p><img src="/img?path=%s" alt="测试图"></p>'
            '</div>') % ip


def test_html_export():
    body = '<h1>演示</h1><p>正文</p>'
    ts = "#preview{color:red}"
    out = md_server.build_standalone_html(None, body, ts)
    check("export-html 非空", bool(out) and len(out) > 200, "len=%d" % len(out))
    check("export-html 含一键复制按钮",
          'id="__copyBtn"' in out and "一键复制" in out)
    check("export-html 含 ClipboardItem 逻辑",
          "ClipboardItem" in out and "navigator.clipboard.write" in out)
    check("export-html 含内联克隆函数", "__inlineClone" in out)
    check("export-html 透传内联样式(#preview->.markdown-body)",
          ".markdown-body{color:red}" in out)
    check("export-html 注入正文", "演示" in out and "正文" in out)


def test_docx_export():
    img = _make_png()
    html = sample_html(img)
    data = docx_export.build_docx(None, html, None)
    check("docx 导出返回非空字节", bool(data) and len(data) > 1000,
          "bytes=%d" % len(data))

    # python-docx 打开校验内容映射
    doc = Document(io.BytesIO(data))
    txt_all = "\n".join(p.text for p in doc.paragraphs)
    check("docx 含标题", any("标题一" in p.text for p in doc.paragraphs))
    check("docx 含列表项", "列表项一" in txt_all and "列表项二" in txt_all)
    check("docx 含引用文字", "引用文字" in txt_all)
    check("docx 含代码块文本", "print(x)" in txt_all)
    check("docx 含表格",
          len(doc.tables) == 1 and doc.tables[0].cell(0, 0).text == "列A")

    # 公式：直接解析 document.xml 校验 m:oMath（python-docx 不会丢外来命名空间）
    dx = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    root = ET.fromstring(dx)
    omaths = root.findall(".//{%s}oMath" % M)
    check("docx 含 m:oMath 公式节点(2)", len(omaths) == 2, "count=%d" % len(omaths))
    check("docx 声明 math 命名空间", M in dx.decode("utf-8", "ignore"))
    for i, om in enumerate(omaths):
        parent = om.getparent()
        check("oMath[%d] 父节点为 w:p" % i,
              parent is not None and parent.tag == "{%s}p" % W)

    # 样式落到 run（红色）
    colored = any(r.font.color is not None and r.font.color.rgb is not None
                  for p in doc.paragraphs for r in p.runs)
    check("docx 文字颜色已应用", colored)

    # 图片嵌入（word/media 存在）
    zf = zipfile.ZipFile(io.BytesIO(data))
    media = [n for n in zf.namelist()
             if n.startswith("word/media/") and not n.endswith("/")]
    check("docx 嵌入图片", len(media) >= 1, "media=%s" % media)


def test_docx_mroot_rad_structure():
    """v2.4.9：带度数根式（\\sqrt[3]）的 OMML 结构——m:deg 必须是
    m:rad 的直接子元素（radPr/deg/e 平级）。历史缺陷：deg 被塞进
    radPr 内属 schema 违规，Word 报"文件损坏"（该分支自 v2.1 写下
    从未被真实用例覆盖，用户文档首次出现 \\sqrt[3] 才暴露）。
    同时覆盖 \\! 剥离重试（v2.4.7）。"""
    BS = chr(92)
    texes = [BS + "sqrt[3]{x}",
             BS + "sqrt[5]{1+" + BS + "sqrt[4]{x}}",
             "f" + BS + "!" + BS + "left(x" + BS + "right)"]
    html = '<div class="markdown-body">'
    for t in texes:
        html += '<div data-tex="%s"></div>' % t
    html += '</div>'
    data = docx_export.build_docx(None, html)
    xml = zipfile.ZipFile(io.BytesIO(data)).read(
        "word/document.xml").decode("utf-8")
    M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    from lxml import etree as _et
    root = _et.fromstring(zipfile.ZipFile(io.BytesIO(data)).read(
        "word/document.xml"))
    rads = root.findall(".//{%s}rad" % M)
    degs = root.findall(".//{%s}deg" % M)
    check("mroot 根式生成 rad 元素", len(rads) >= 3, "rad=%d" % len(rads))
    check("deg 平级于 rad（radPr/deg/e 三层结构）", len(degs) >= 3,
          "deg=%d" % len(degs))
    bad = 0
    for rp in root.findall(".//{%s}radPr" % M):
        for c in rp:
            if isinstance(c.tag, str) and _et.QName(c).localname != "degHide":
                bad += 1
    check("radPr 内无 deg（schema 合规）", bad == 0, "bad=%d" % bad)
    for rad in rads:
        kids = [_et.QName(c).localname for c in rad
                if isinstance(c.tag, str)]
        if kids and kids[0] == "radPr":
            # deg 在 e 之前
            if "deg" in kids and kids.index("deg") > kids.index("e"):
                bad += 1
    check("deg 在 e 之前（rad 子元素顺序）", bad == 0)
    check("\\! 公式经剥离重试仍出 OMML",
          xml.count("<m:oMath>") == 3, "oMath=%d" % xml.count("<m:oMath>"))


def test_docx_pblk_export():
    """v1.11.5 场景：前端 inlineStylesInto 之后的真实结构——
    #preview 克隆为裸 div（id/class 已剥），pblk 为裸 div，图表为栅格化
    data URL PNG 或回落 SVG。历史缺陷：裸 div 被当段落拉平（标题/列表/
    表格语义全灭），SVG 文本被抽成 CSS 垃圾文字。"""
    png1px = ("data:image/png;base64,"
              "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
              "AAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
    body = (
        '<div>'
        '<div>'
        '<h1>接力标题</h1>'
        '<p style="color:#cc0000;font-size:18px;text-align:center">'
        '居中红字 <strong>加粗</strong></p>'
        '<ul><li>接力列表</li></ul>'
        '<table><tr><th>T</th></tr><tr><td>v</td></tr></table>'
        '</div>'
        '<div style="text-align:center"><img src="%s"/></div>'
        '<div><svg xmlns="http://www.w3.org/2000/svg" width="10" height="10">'
        '<style>#mmd-9{font-family:trebuchet ms;fill:#ccc;}</style>'
        '<text>lbl</text></svg></div>'
        '<div data-tex="E = mc^2"></div>'
        '<div>纯文本兜底</div>'
        '</div>'
    ) % png1px
    data = docx_export.build_docx(None, body)
    doc = Document(io.BytesIO(data))
    paras = doc.paragraphs
    alltext = "\n".join(p.text for p in paras)
    styles = [p.style.name for p in paras]
    dxml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    check("pblk 裸 div 下钻：H1 保留 Heading 1",
          any(p.style.name == "Heading 1" and p.text == "接力标题" for p in paras))
    check("pblk 裸 div 下钻：列表保留 List Bullet", "List Bullet" in styles)
    check("pblk 裸 div 下钻：表格保留", len(doc.tables) == 1
          and doc.tables[0].cell(0, 0).text == "T")
    red = next((r for p in paras for r in p.runs if r.text.startswith("居中红字")), None)
    check("pblk 裸 div 下钻：颜色生效", red is not None
          and str(red.font.color.rgb) == "CC0000")
    check("pblk 裸 div 下钻：字号生效(18px->13.5pt)", red is not None
          and red.font.size is not None and abs(red.font.size.pt - 13.5) < 0.01)
    check("data URL 图表 PNG 嵌入", len(doc.inline_shapes) == 1,
          "shapes=%d" % len(doc.inline_shapes))
    check("SVG 回落：占位 [图表] 无 CSS 垃圾",
          "[图表]" in alltext and "mmd" not in alltext
          and "trebuchet" not in alltext)
    check("data-tex 公式仍出 OMML", b"oMath" in dxml)
    check("裸 div 直接文本兜底不丢", "纯文本兜底" in alltext)


def test_docx_img_width_rules():
    """v2.1.2/v2.2.0 图片宽度规则（与前端 {N} 同口径，默认 A4
    版心 174mm=6.8504in）：data-x-pct=30 → 版心×0.3；data-x-w=227px
    → 227/96 钳版心；自然尺寸：大图 800px 钳版心、小图 300px 原始"""
    import base64
    img = _make_png()   # 8x8 小图
    from PIL import Image
    big = img.replace("px.png", "big.png")
    Image.new("RGB", (800, 200), (60, 60, 200)).save(big)
    small = img.replace("px.png", "small.png")
    Image.new("RGB", (300, 100), (60, 200, 60)).save(small)
    body = (
        '<div class="markdown-body">'
        '<p><img src="data:image/png;base64,%s" data-x-pct="30" '
        'style="width:30%%"/></p>'
        '<p><img src="data:image/png;base64,%s" data-x-w="227" '
        'style="width:227px;max-width:100%%;height:auto"/></p>'
        '<p><img src="/img?path=%s"/></p>'
        '<p><img src="/img?path=%s"/></p>'
        '<p><img src="/img?path=%s" data-x-pct="50" style="width:50%%"/></p>'
        '</div>'
    ) % (
        base64.b64encode(open(small, "rb").read()).decode(),
        base64.b64encode(open(small, "rb").read()).decode(),
        big.replace("\\", "/"), small.replace("\\", "/"),
        big.replace("\\", "/"),
    )
    data = docx_export.build_docx(None, body)
    doc = Document(io.BytesIO(data))
    shapes = doc.inline_shapes
    check("宽度规则：图片全部嵌入", len(shapes) == 5,
          "shapes=%d" % len(shapes))
    if len(shapes) == 5:
        USABLE = 174 / 25.4   # A4 版心 210−18−18=174mm
        check("pct=30 → 版心×0.3（2.055in）",
              abs(shapes[0].width.inches - USABLE * 0.3) < 0.01,
              "%.3f" % shapes[0].width.inches)
        check("data-x-w=227px → 2.364in（< 版心不钳）",
              abs(shapes[1].width.inches - 227 / 96.0) < 0.01,
              "%.3f" % shapes[1].width.inches)
        check("大图 800px 钳版心 6.85in",
              abs(shapes[2].width.inches - USABLE) < 0.01,
              "%.3f" % shapes[2].width.inches)
        check("小图 300px 原始 3.125in",
              abs(shapes[3].width.inches - 300 / 96.0) < 0.01,
              "%.3f" % shapes[3].width.inches)
        check("pct=50 → 版心×0.5（3.425in）",
              abs(shapes[4].width.inches - USABLE * 0.5) < 0.01,
              "%.3f" % shapes[4].width.inches)


def test_regression_import():
    # 轻量回归：md_server 正常导入（既有路由无语法/导入错误），且关键函数仍在
    check("md_server 导入且 _rewrite_assets_ref 存在",
          hasattr(md_server, "_rewrite_assets_ref"))
    check("md_server build_standalone_html 存在",
          hasattr(md_server, "build_standalone_html"))


def main():
    test_html_export()
    test_docx_export()
    test_docx_pblk_export()
    test_docx_img_width_rules()
    test_docx_mroot_rad_structure()
    test_regression_import()
    n_pass = sum(1 for _, ok in RESULTS if ok)
    n_all = len(RESULTS)
    print("%d/%d passed" % (n_pass, n_all))
    return 0 if n_pass == n_all else 1


if __name__ == "__main__":
    sys.exit(main())
